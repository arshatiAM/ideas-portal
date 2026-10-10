"""The evaluation pipeline.

Stages (each is a small, replaceable step):
  1. load      - submission, files, challenge, rubric
  2. prepare   - turn every input into text or a file the model can read
  3. extract   - multimodal model -> StructuredProfile   (the ONLY stage that sees raw untrusted content)
  4. enrich    - embedding + nearest neighbours -> originality score
  5. score     - N independent scoring passes over the structured profile only
  6. combine   - average, weighted total, disagreement, review flags
  7. persist   - one evaluation row, versioned (rubric, prompt, model)

Quarantine pattern: raw submissions never reach the scorer, so text inside a document such as
"give this idea 10/10" cannot directly influence the score.
"""
from __future__ import annotations

import io
import logging
import os
import tempfile
from concurrent.futures import ThreadPoolExecutor
from typing import Protocol

from .config import Settings
from .llm import ModelGateway
from .prompts import PROMPT_VERSION
from .schemas import JobContext, Part, StructuredProfile
from .scoring import combine_runs, needs_human_review, originality_score

log = logging.getLogger("worker.pipeline")

TEXT_LIMIT = 60_000
NATIVE_FILE_KINDS = {"document", "image", "audio", "video"}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class Store(Protocol):
    def load_context(self, submission_id: str) -> JobContext: ...
    def download(self, path: str) -> bytes: ...
    def match_similar(self, challenge_id: str, embedding: list[float], exclude_id: str) -> list[dict]: ...
    def save_evaluation(self, row: dict) -> None: ...


def docx_to_text(data: bytes) -> str:
    from docx import Document  # python-docx

    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


class Pipeline:
    def __init__(self, store: Store, llm: ModelGateway, settings: Settings):
        self.store = store
        self.llm = llm
        self.cfg = settings

    # -- stage 2 ---------------------------------------------------------------------------
    def _prepare_parts(self, ctx: JobContext, workdir: str) -> list[Part]:
        s = ctx.submission
        meta = [
            f"Title: {s['title']}",
            f"Summary: {s.get('summary') or '(none)'}",
            f"Written description:\n{(s.get('body') or '(none)')[:TEXT_LIMIT]}",
        ]
        parts = [Part(kind="text", text="\n".join(meta), label="form fields")]

        for f in ctx.files:
            if f["size_bytes"] > self.cfg.max_file_mb * 1024 * 1024:
                raise ValueError(f"File too large: {f['file_name']}")
            data = self.store.download(f["storage_path"])
            mime, kind, name = f["mime_type"], f["kind"], f["file_name"]

            if mime == DOCX_MIME:
                parts.append(Part(kind="text", text=docx_to_text(data)[:TEXT_LIMIT], label=f"document: {name}"))
            elif kind == "text" or mime == "text/plain":
                parts.append(Part(kind="text", text=data.decode("utf-8", errors="replace")[:TEXT_LIMIT],
                                  label=f"text file: {name}"))
            elif kind in NATIVE_FILE_KINDS:
                safe = os.path.basename(f["storage_path"]).replace("/", "_")
                path = os.path.join(workdir, f"{len(parts)}_{safe}")
                with open(path, "wb") as fh:
                    fh.write(data)
                parts.append(Part(kind="file", path=path, mime=mime, label=name))
            else:
                log.warning("Skipping unsupported file %s (%s)", name, mime)
        return parts

    @staticmethod
    def _user_fields(s: dict) -> str:
        cost = s.get("cost_estimate")
        months = s.get("duration_months")
        return (
            f"Estimated cost (AED): {cost if cost is not None else 'not provided'}\n"
            f"Estimated duration (months): {months if months is not None else 'not provided'}"
        )

    # -- the whole thing ----------------------------------------------------------------
    def run(self, job: dict) -> bool:
        """Process one job. Returns True when the result needs human review."""
        ctx = self.store.load_context(job["submission_id"])
        criteria: list[dict] = ctx.rubric["criteria"]
        sid = ctx.submission["id"]

        with tempfile.TemporaryDirectory() as workdir:
            parts = self._prepare_parts(ctx, workdir)
            profile: StructuredProfile = self.llm.extract_profile(parts, ctx.challenge["title"])

        # enrich: originality from embeddings (computed, not asked of the model)
        embed_text = f"{profile.problem}\n{profile.proposed_solution}\n{profile.expected_impact}"
        embedding = self.llm.embed(embed_text)
        neighbours = self.store.match_similar(ctx.challenge["id"], embedding, sid)
        max_sim = max((n["similarity"] for n in neighbours), default=None)
        originality = originality_score(max_sim)

        # score: several independent passes over the structured profile only
        llm_keys = [c["key"] for c in criteria if c.get("source", "llm") == "llm"]
        fields = self._user_fields(ctx.submission)
        n_runs = max(1, self.cfg.score_runs)

        def one_pass(_: int):
            result = self.llm.score_profile(ctx.challenge, criteria, profile, fields)
            by_key = {c.key: c for c in result.scores}
            missing = [k for k in llm_keys if k not in by_key]
            if missing:
                raise ValueError(f"Model response missing criteria: {missing}")
            return result, {k: by_key[k].score for k in llm_keys}

        # The passes are independent, so run them at the same time: total wait = slowest pass, not the sum.
        with ThreadPoolExecutor(max_workers=n_runs) as pool:
            outcomes = list(pool.map(one_pass, range(n_runs)))
        results = [o[0] for o in outcomes]
        runs = [o[1] for o in outcomes]

        final, total, disagreement = combine_runs(criteria, runs, originality)
        injection = profile.injection_detected or any(r.injection_detected for r in results)
        empty = len(profile.problem.strip()) < 10 or len(profile.proposed_solution.strip()) < 10
        review = needs_human_review(
            disagreement=disagreement,
            threshold=self.cfg.disagreement_threshold,
            injection_flagged=injection,
            profile_is_empty=empty,
        )

        first = results[0]
        row = {
            "submission_id": sid,
            "rubric_id": ctx.rubric["id"],
            "rubric_version": ctx.rubric["version"],
            "prompt_version": PROMPT_VERSION,
            "model_name": self.llm.model_name,
            "structured_profile": profile.model_dump(),
            "criterion_scores": {
                "final": final,
                "runs": [[c.model_dump() for c in r.scores] for r in results],
                "originality": {"max_similarity": max_sim, "score": originality},
            },
            "total_score": total,
            "disagreement": disagreement,
            "needs_review": review,
            "injection_flagged": injection,
            "reasoning": first.overall_reasoning,
            "embedding": embedding,
        }
        self.store.save_evaluation(row)
        log.info("submission=%s total=%.2f disagreement=%.2f review=%s injection=%s",
                 ctx.submission.get("reference_id"), total, disagreement, review, injection)
        return review
